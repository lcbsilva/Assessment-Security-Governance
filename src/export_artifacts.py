#!/usr/bin/env python3
"""Gera XLSX, PPTX e PDF resumidos a partir do contrato normalizado.

Os formatos são complementares ao HTML. O conteúdo é derivado do JSON local;
nenhuma chamada externa é feita por este script.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from insight_engine import prioritize_findings
from executive_intelligence import build as build_executive_intelligence
from local_privacy import protect_output_directory
from report_context import build as build_report_context
from quality_audit import audit


def main() -> None:
    parser = argparse.ArgumentParser(description="Exporta artefatos do assessment")
    parser.add_argument("--data", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, default=Path("dist"))
    args = parser.parse_args()
    data = json.loads(args.data.read_text(encoding="utf-8"))
    protect_output_directory(args.output_dir)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_xlsx(data, args.output_dir / "assessment-action-plan.xlsx")
    write_pptx(data, args.output_dir / "assessment-executive-summary.pptx")
    write_pdf(data, args.output_dir / "assessment-executive-summary.pdf")
    catalog = __import__("yaml").safe_load((Path(__file__).resolve().parents[1] / "catalog" / "controls.yaml").read_text(encoding="utf-8"))
    write_one_page_brief(data, args.output_dir / "assessment-one-page-brief.pdf", catalog)
    print(f"Artefatos exportados em {args.output_dir}")


def write_xlsx(data: dict, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    context = build_report_context(data)
    cover = book.active
    cover.title = "Resumo e confidencialidade"
    cover.append(["Security & Governance Assessment", "Classificação e contexto da execução"])
    cover.append(["Classificação", context["classification"]])
    cover.append(["Manuseio", "Confidencial · compartilhar somente com pessoas autorizadas"])
    cover.append(["Execução read-only", "Sim" if data.get("metadata", {}).get("execution", {}).get("tenant_mutation") is False else "Não confirmado"])
    cover.append(["Perfil", context["profile"]])
    cover.append(["Run ID", context["run_id"]])
    cover.append(["Versão do engine", context["engine_version"]])
    cover.append(["Início UTC", context["started_at"]])
    cover.append(["Fim UTC", context["finished_at"]])
    cover.append(["Status do contrato", context["contract_status"]])
    cover.append([])
    cover.append(["Escopo observado", "Quantidade"])
    for row in context["scope_rows"]:
        cover.append([row["label"], row["value"]])
    cover.append([])
    cover.append(["Limitações registradas", len(context["limitations"])])
    cover.append(["Aviso", "Contagens de objetos não comprovam impacto de negócio, incidente ou indisponibilidade."])
    cover.append([])
    cover.append(["Módulos com limitação", "Estado · quantidade · interpretação"])
    for item in context["limitations"]:
        cover.append([item["module"], f"{item['status']} · {item['records']} · provável: {item['likely_cause']} Próximo passo: {item['next_step']}"])
    for column in cover.columns:
        cover.column_dimensions[column[0].column_letter].width = 42
    for cell in cover[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B2C83")
    cover.freeze_panes = "A2"
    sheet = book.create_sheet("Plano de ação", 1)
    headers = ["ID", "Achado", "Severidade", "Impacto técnico", "Risco", "Alcance observado", "Impacto potencial em usuários", "Prioridade", "Score prioridade", "Confiança por achado", "Fonte da evidência", "Módulo", "Estado da coleta", "Janela de coleta", "Sinal financeiro", "Esforço relativo", "Quadrante", "Frente consultiva", "Responsável", "Status", "Dependências", "30 dias", "60 dias", "90 dias"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B2C83")
    cost_signal = data.get("discovery", {}).get("lifecycle", {}).get("summary", {}).get("Custo mensal potencial", "Não quantificado")
    findings = prioritize_findings(data.get("findings", []), data.get("metadata", {}).get("evidence_quality", {}), cost_signal)
    executive = build_executive_intelligence(findings)
    for item in findings:
        actions = item.get("action_30_60_90", {})
        severity = item.get("severity", "medium")
        impact = "Muito alto" if severity == "critical" else ("Alto" if severity == "high" else ("Médio" if severity == "medium" else "Baixo"))
        effort = int(item.get("effort", 3) or 3)
        risk = int(item.get("risk_score", 0) or 0)
        quadrant = "Alto impacto / baixo esforço" if risk >= 70 and effort <= 2 else ("Alto impacto / alto esforço" if risk >= 70 else ("Baixo impacto / baixo esforço" if effort <= 2 else "Baixo impacto / alto esforço"))
        workstreams = {"ID": "Identity & Access", "SEC": "Cloud & M365 Security", "GOV": "Cloud Governance", "COST": "FinOps & Cloud Optimization"}
        workstream = workstreams.get(str(item.get("control_id", "")).split("-")[0], "Risk & Compliance")
        dependencies = "; ".join(item.get("remediation_dependencies", []))
        lineage = item.get("evidence_lineage", {})
        window = lineage.get("collection_window", {}) or {}
        collection_window = f"{window.get('started_at') or 'N/D'} → {window.get('finished_at') or 'N/D'}"
        sheet.append([item.get("id"), item.get("title"), severity, impact, risk, f"{item.get('affected', 'N/D')} {item.get('affected_unit', 'itens')}", item.get("user_impact", "Não determinado; validar com o owner."), item.get("priority", "P3"), item.get("priority_score", 0), item.get("evidence_confidence", "baixa"), lineage.get("source", item.get("source", "Não informado")), lineage.get("module", "Não vinculado"), lineage.get("source_status", "Não informado"), collection_window, item.get("financial_signal", "unquantified"), item.get("effort_band", effort), quadrant, workstream, item.get("owner"), item.get("status"), dependencies, actions.get("30"), actions.get("60"), actions.get("90")])
    for insight in data.get("discovery", {}).get("cross_domain_insights", []):
        risk = int(insight.get("risk", 0) or 0)
        owner = insight.get("suggested_owner", "Security & Governance")
        sheet.append([insight.get("id"), insight.get("title"), insight.get("severity", "medium"), "Alto" if risk >= 75 else "Médio", risk, f"{insight.get('affected', 'N/D')} itens", "Não determinado; validar com o owner.", insight.get("priority", "P3"), risk, "média", "Evidências cruzadas", "Múltiplos módulos", "Requer validação", "N/D", "unquantified", insight.get("effort_band", "Médio"), "Alto impacto / validar esforço" if risk >= 75 else "Revisão planejada", owner, owner, "Insight — validação", "Validar evidência, escopo e owner", "Executar somente o plano aprovado", "Reavaliar evidência e registrar evolução"])
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
    sheet.freeze_panes = "A2"

    exec_sheet = book.create_sheet("Prioridades executivas", 2)
    summary = executive["summary"]
    exec_sheet.append(["Resumo executivo", "Valor"])
    exec_sheet.append(["Achados priorizados", summary["findings"]])
    exec_sheet.append(["P1", summary["p1"]])
    exec_sheet.append(["P2", summary["p2"]])
    exec_sheet.append(["P3", summary["p3"]])
    exec_sheet.append(["Confirmados para ação", summary["confirmed_for_action"]])
    exec_sheet.append(["Revisão condicional", summary["conditional_review"]])
    exec_sheet.append(["Quick wins", summary["quick_wins"]])
    exec_sheet.append([])
    exec_sheet.append(["Quick wins", "Prioridade", "Risco", "Esforço", "Owner", "Resultado esperado"])
    for item in executive["quick_wins"]:
        exec_sheet.append([item["title"], item["priority"], item["risk"], item["effort"], item["owner"], item["outcome"]])
    exec_sheet.append([])
    exec_sheet.append(["Frente consultiva", "Achados", "Confirmados", "Condicionais", "Maior risco", "Prioridade"])
    for item in executive["workstreams"]:
        exec_sheet.append([item["name"], item["findings"], item["confirmed"], item["conditional"], item["max_risk"], item["top_priority"]])
    exec_sheet.append([])
    exec_sheet.append(["Horizonte", "Prioridade", "Achado", "Owner", "Ação"])
    for horizon, actions in executive["roadmap"].items():
        for item in actions:
            exec_sheet.append([f"{horizon} dias", item["priority"], item["finding"], item["owner"], item["action"]])
    for cell in exec_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B2C83")
    for column in exec_sheet.columns:
        exec_sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 64)
    exec_sheet.freeze_panes = "A2"

    coverage_sheet = book.create_sheet("Cobertura e limitações", 3)
    coverage_headers = ["Módulo", "Domínio", "Escopo esperado", "Status", "Registros", "Confiança", "Classificação", "Causa provável", "Próximo passo", "Observação técnica"]
    coverage_sheet.append(coverage_headers)
    for cell in coverage_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B2C83")
    coverage_rows = data.get("metadata", {}).get("coverage_map", []) or []
    if coverage_rows:
        for item in coverage_rows:
            coverage_sheet.append([
                item.get("module"), item.get("domain"), item.get("expected_read_scope"), item.get("status"),
                item.get("records", 0), item.get("evidence_confidence"), item.get("limitation_category", item.get("category", "unknown")),
                item.get("likely_cause", "—"), item.get("next_step", "—"), item.get("limitation", "—"),
            ])
    else:
        for item in context["limitations"]:
            coverage_sheet.append([
                item.get("module"), "N/D", "N/D", item.get("status"), item.get("records", 0), "baixa",
                item.get("limitation_category"), item.get("likely_cause"), item.get("next_step"), item.get("summary"),
            ])
    for column in coverage_sheet.columns:
        coverage_sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
    coverage_sheet.freeze_panes = "A2"
    coverage_sheet.auto_filter.ref = coverage_sheet.dimensions

    external = data.get("discovery", {}).get("external_assessments", {}).get("microsoft_zero_trust")
    if external:
        def safe_cell(value: object) -> object:
            """Evita que texto de arquivo importado seja interpretado como fórmula Excel."""
            if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
                return "'" + value
            return value

        external_sheet = book.create_sheet("Zero Trust Microsoft")
        external_sheet.append(["Fonte", external.get("source", "Microsoft Zero Trust Assessment")])
        external_sheet.append(["Versão", external.get("source_version", "N/D")])
        external_sheet.append(["Importado como referência; não combinado ao score nativo", "Sim"])
        external_sheet.append([])
        headers = ["ID Microsoft", "Pilar(es)", "Verificação", "Estado Microsoft", "Risco do teste", "Licença mínima", "Controles relacionados"]
        external_sheet.append(headers)
        for cell in external_sheet[5]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="5B2C83")
        for item in external.get("tests", []):
            external_sheet.append([
                safe_cell(item.get("test_id")), safe_cell(", ".join(item.get("pillars", []))), safe_cell(item.get("title")),
                safe_cell(item.get("status")), safe_cell(item.get("risk")), safe_cell(", ".join(item.get("minimum_license", [])) or "Não indicado"),
                safe_cell(", ".join(item.get("mapped_control_ids", [])) or "Sem correspondência direta"),
            ])
        for column in external_sheet.columns:
            external_sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 64)
        external_sheet.freeze_panes = "A6"

    intelligence = data.get("discovery", {})
    intel_sheet = book.create_sheet("Azure Intelligence")
    intel_sheet.append(["Resource Intelligence", "Valor"])
    for cell in intel_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B2C83")
    graph = intelligence.get("resource_map", {})
    hygiene = intelligence.get("resource_hygiene", {})
    secure = intelligence.get("defender_secure_score", {})
    intel_sheet.append(["Recursos no mapa", graph.get("node_count", 0)])
    intel_sheet.append(["Relações demonstradas", graph.get("edge_count", 0)])
    intel_sheet.append(["Resource groups vazios", hygiene.get("empty_resource_group_count", 0)])
    intel_sheet.append(["Recursos órfãos/não associados", hygiene.get("orphan_count", 0)])
    intel_sheet.append(["Sinais de rede em atenção", hygiene.get("network_attention_count", 0)])
    intel_sheet.append([])
    intel_sheet.append(["Secure Score / Subscription", "Atual", "Máximo", "Percentual"])
    for item in secure.get("scores", []):
        intel_sheet.append([item.get("subscription"), item.get("current"), item.get("max"), item.get("percentage")])
    intel_sheet.append([])
    intel_sheet.append(["Controle Secure Score", "Atual", "Máximo", "Ganho potencial", "Recursos não saudáveis"])
    for item in secure.get("top_improvements", []):
        intel_sheet.append([item.get("control"), item.get("score"), item.get("max_score"), item.get("potential_score_increase"), item.get("unhealthy_resources")])
    intel_sheet.append([])
    intel_sheet.append(["Policy Assignment", "Escopo", "Definition", "Enforcement", "Parâmetro", "Default", "Assigned", "Effective", "Origem"])
    for assignment in intelligence.get("policy_assignments", []):
        if assignment.get("parameters"):
            for parameter in assignment.get("parameters", []):
                intel_sheet.append([assignment.get("assignment"), assignment.get("scope"), assignment.get("definition_display_name"), assignment.get("enforcement_mode"), parameter.get("name"), parameter.get("default_value"), parameter.get("assigned_value"), parameter.get("effective_value"), parameter.get("value_source")])
        else:
            intel_sheet.append([assignment.get("assignment"), assignment.get("scope"), assignment.get("definition_display_name"), assignment.get("enforcement_mode"), "—", "—", "—", "—", "Not set"])
    for column in intel_sheet.columns:
        intel_sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 55)
    intel_sheet.freeze_panes = "A2"
    book.save(path)


def write_pptx(data: dict, path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    meta = data.get("metadata", {})
    engagement = meta.get("engagement", {}) or {}
    customer = meta.get("customer_name") or engagement.get("customer_name", "Tenant")
    engagement_name = meta.get("engagement_name") or engagement.get("engagement_name", "Security & Governance Assessment")
    consultant = meta.get("consultant_name") or engagement.get("consultant_name", "Consultor não informado")
    classification = meta.get("classification") or engagement.get("classification", "Confidencial — Security & Governance Assessment")
    findings = prioritize_findings(data.get("findings", []), data.get("metadata", {}).get("evidence_quality", {}), data.get("discovery", {}).get("lifecycle", {}).get("summary", {}).get("Custo mensal potencial", "Não quantificado"))
    executive = build_executive_intelligence(findings)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.8), Inches(12), Inches(1.2))
    box.text_frame.text = engagement_name
    box.text_frame.paragraphs[0].font.size = Pt(28)
    box.text_frame.paragraphs[0].font.bold = True
    sub = slide.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(11), Inches(0.8))
    context = build_report_context(data)
    meta = data.get("metadata", {})
    profile = context["profile"]
    if profile in {"não informado", "unknown", ""} and meta.get("simulation", {}).get("is_simulation") is True:
        profile = "Demonstração sintética"
    scope_text = " · ".join(f"{row['label']}: {row['value']}" for row in context["scope_rows"][:4])
    sub.text_frame.text = f"{customer} · {consultant} · {classification} · Execução {context['started_at']} · Perfil {profile}\n{scope_text}\nConfidencial · compartilhar somente com pessoas autorizadas"
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11), Inches(0.6))
    title.text_frame.text = "Riscos prioritários"
    title.text_frame.paragraphs[0].font.size = Pt(24)
    text = slide.shapes.add_textbox(Inches(0.9), Inches(1.4), Inches(11), Inches(5.2))
    text.text_frame.text = "\n".join(f"• {item.get('title')} — risco {item.get('risk_score')}/100 · {item.get('affected', 'N/D')} {item.get('affected_unit', 'itens')} · esforço {item.get('effort_band', 'não avaliado')} · confiança {item.get('evidence_confidence', 'não avaliada')} · fonte {item.get('evidence_lineage', {}).get('source', item.get('source', 'N/D'))} ({item.get('evidence_lineage', {}).get('source_status', 'estado N/D')})" for item in findings[:5]) or "Nenhum achado disponível"
    insights = data.get("discovery", {}).get("cross_domain_insights", [])
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11), Inches(0.6))
    title.text_frame.text = "Insights cruzados — Segurança & Governança"
    title.text_frame.paragraphs[0].font.size = Pt(24)
    text = slide.shapes.add_textbox(Inches(0.9), Inches(1.4), Inches(11), Inches(5.2))
    text.text_frame.text = "\n".join(f"• {item.get('priority', 'P3')} · {item.get('title')} — {item.get('suggested_owner', 'Security & Governance')} · {item.get('effort_band', 'Médio')}" for item in insights[:5]) or "Nenhum insight cruzado disponível"
    context = build_report_context(data)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11.5), Inches(0.7))
    title.text_frame.text = "Cobertura e limitações da execução"
    title.text_frame.paragraphs[0].font.size = Pt(22)
    status_text = " · ".join(f"{status}: {count}" for status, count in context["status_counts"].items()) or "Sem manifesto de coletores"
    limitation_lines = [f"• {item['module']} — {item['status']} ({item['limitation_category']}): {item['likely_cause']} Próximo passo: {item['next_step']}" for item in context["limitations"][:8]]
    body = slide.shapes.add_textbox(Inches(0.9), Inches(1.4), Inches(11.2), Inches(4.8))
    body.text_frame.text = f"Perfil {context['profile']} · contrato {context['contract_status']}\nMódulos: {status_text}\n\n" + ("\n".join(limitation_lines) if limitation_lines else "Nenhuma limitação de módulo registrada.") + "\n\nContagens de alcance não comprovam impacto operacional; validar com os owners."
    body.text_frame.paragraphs[0].font.size = Pt(14)

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11.5), Inches(0.7))
    title.text_frame.text = "Decisão executiva — próximos passos"
    title.text_frame.paragraphs[0].font.size = Pt(22)
    actionable = [
        item for item in context["limitations"]
        if item.get("status") in {"partial", "not_available", "error"}
    ][:5]
    action_lines = [
        f"• {item['module']} — {item['limitation_category']}: {item['next_step']}"
        for item in actionable
    ]
    top_risks = [
        f"• {item.get('priority', 'P3')} · {item.get('title')} — owner {item.get('owner', 'A definir')}"
        for item in findings[:3]
    ]
    quick_win_lines = [
        f"• {item['priority']} · {item['title']} — risco {item['risk']}/100 · esforço {item['effort']} · owner {item['owner']}"
        for item in executive["quick_wins"][:3]
    ]
    summary = executive["summary"]
    body = slide.shapes.add_textbox(Inches(0.9), Inches(1.3), Inches(11.2), Inches(5.2))
    body.text_frame.text = (
        f"Decisão: {summary['p1']} P1 · {summary['p2']} P2 · {summary['conditional_review']} dependem de validação de evidência.\n\n"
        "Quick wins confirmados:\n" + ("\n".join(quick_win_lines) if quick_win_lines else "• Nenhum quick win confirmado com evidência suficiente.") +
        "\n\nPrioridades de risco:\n" + ("\n".join(top_risks) if top_risks else "• Nenhum achado priorizado nesta execução.") +
        "\n\nPendências de coleta:\n" + ("\n".join(action_lines[:3]) if action_lines else "• Nenhuma pendência operacional crítica registrada.") +
        "\n\nPrincípio: ausência de evidência não é conformidade; validar owner, escopo e dependências antes de remediar."
    )
    body.text_frame.paragraphs[0].font.size = Pt(14)

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11.5), Inches(0.7))
    title.text_frame.text = "Roadmap executivo — 30 / 60 / 90 dias"
    title.text_frame.paragraphs[0].font.size = Pt(22)
    roadmap_lines = []
    for horizon in ("30", "60", "90"):
        actions = executive["roadmap"][horizon][:3]
        roadmap_lines.append(f"{horizon} dias")
        roadmap_lines.extend(
            f"• {item['priority']} · {item['finding']} — {item['action']} · owner {item['owner']}"
            for item in actions
        )
        if not actions:
            roadmap_lines.append("• Nenhuma ação confirmada neste horizonte.")
    body = slide.shapes.add_textbox(Inches(0.9), Inches(1.3), Inches(11.2), Inches(5.4))
    body.text_frame.text = "\n".join(roadmap_lines) + "\n\nItens com evidência insuficiente permanecem fora do roadmap confirmado até validação."
    body.text_frame.paragraphs[0].font.size = Pt(13)

    external = data.get("discovery", {}).get("external_assessments", {}).get("microsoft_zero_trust")
    if external:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11.5), Inches(0.7))
        title.text_frame.text = "Microsoft Zero Trust Assessment — referência externa"
        title.text_frame.paragraphs[0].font.size = Pt(22)
        summary = external.get("status_counts", {})
        status_line = " · ".join(f"{key}: {value}" for key, value in summary.items())
        pillar_lines = []
        for pillar in external.get("pillars", []):
            pillar_lines.append(f"{pillar.get('name')}: {pillar.get('tests')} verificações · " + ", ".join(f"{key} {value}" for key, value in pillar.get("statuses", {}).items()))
        body = slide.shapes.add_textbox(Inches(0.9), Inches(1.4), Inches(11.2), Inches(4.8))
        body.text_frame.text = f"Versão {external.get('source_version', 'N/D')} · {external.get('tests_total', 0)} verificações\n{status_line}\n\n" + "\n".join(pillar_lines) + "\n\nScore Microsoft mantido separado; não combinado ao score do Assessment Engine."
        body.text_frame.paragraphs[0].font.size = Pt(14)

    discovery = data.get("discovery", {})
    graph = discovery.get("resource_map", {})
    hygiene = discovery.get("resource_hygiene", {})
    secure = discovery.get("defender_secure_score", {})
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11.5), Inches(0.7))
    title.text_frame.text = "Azure Intelligence — mapa, governança e Secure Score"
    title.text_frame.paragraphs[0].font.size = Pt(22)
    top_controls = secure.get("top_improvements", [])[:3]
    controls_text = "\n".join(f"• {item.get('control')} — ganho potencial {item.get('potential_score_increase')} ponto(s)" for item in top_controls) or "• Secure Score não disponível"
    body = slide.shapes.add_textbox(Inches(0.9), Inches(1.4), Inches(11.2), Inches(4.8))
    body.text_frame.text = (
        f"Resource Map: {graph.get('node_count', 0)} recursos · {graph.get('edge_count', 0)} relações demonstradas\n"
        f"Higiene: {hygiene.get('empty_resource_group_count', 0)} RGs vazios · {hygiene.get('orphan_count', 0)} órfãos/não associados · {hygiene.get('network_attention_count', 0)} sinais de rede\n"
        f"Policy assignments: {len(discovery.get('policy_assignments', []))}\n\n"
        f"Controles com maior potencial de melhoria:\n{controls_text}\n\n"
        "Relações ausentes e defaults não retornados permanecem como evidência insuficiente; nenhuma remediação é executada."
    )
    body.text_frame.paragraphs[0].font.size = Pt(14)
    presentation.save(path)


def write_pdf(data: dict, path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    document = SimpleDocTemplate(str(path), pagesize=A4, title="Security & Governance Assessment")
    meta = data.get("metadata", {})
    context = build_report_context(data)
    engagement = meta.get("engagement", {}) or {}
    customer = meta.get("customer_name") or engagement.get("customer_name", "Tenant")
    engagement_name = meta.get("engagement_name") or engagement.get("engagement_name", "Security & Governance Assessment")
    classification = meta.get("classification") or engagement.get("classification", "Confidencial — Security & Governance Assessment")
    findings = prioritize_findings(data.get("findings", []), data.get("metadata", {}).get("evidence_quality", {}), data.get("discovery", {}).get("lifecycle", {}).get("summary", {}).get("Custo mensal potencial", "Não quantificado"))
    executive = build_executive_intelligence(findings)
    scope_line = " · ".join(f"{row['label']}: {row['value']}" for row in context["scope_rows"])
    story = [Paragraph(str(engagement_name), styles["Title"]), Paragraph(str(customer), styles["Heading2"]), Paragraph(str(classification), styles["BodyText"]), Paragraph("Confidencial · compartilhar somente com pessoas autorizadas", styles["BodyText"]), Paragraph(f"Perfil {context['profile']} · Run ID {context['run_id']} · Início UTC {context['started_at']} · fim UTC {context['finished_at']}", styles["BodyText"]), Paragraph(scope_line, styles["BodyText"]), Spacer(1, 18)]
    story.append(Paragraph("Top riscos", styles["Heading2"]))
    for item in findings[:5]:
        lineage = item.get("evidence_lineage", {})
        story.append(Paragraph(f"{item.get('title')} — risco {item.get('risk_score')}/100 · alcance: {item.get('affected', 'N/D')} {item.get('affected_unit', 'itens')} · esforço: {item.get('effort_band', 'não avaliado')} · confiança: {item.get('evidence_confidence', 'não avaliada')} · fonte: {lineage.get('source', item.get('source', 'N/D'))} ({lineage.get('source_status', 'estado N/D')})", styles["BodyText"]))
        story.append(Spacer(1, 6))
    story.append(Paragraph("Decisão executiva", styles["Heading2"]))
    summary = executive["summary"]
    story.append(Paragraph(f"P1: {summary['p1']} · P2: {summary['p2']} · quick wins confirmados: {summary['quick_wins']} · revisões condicionais: {summary['conditional_review']}.", styles["BodyText"]))
    for item in executive["quick_wins"][:5]:
        story.append(Paragraph(f"Quick win — {item['priority']} · {item['title']} · risco {item['risk']}/100 · esforço {item['effort']} · owner {item['owner']}", styles["BodyText"]))
    story.append(Paragraph("Roadmap 30 / 60 / 90", styles["Heading2"]))
    for horizon in ("30", "60", "90"):
        actions = executive["roadmap"][horizon][:4]
        if actions:
            for item in actions:
                story.append(Paragraph(f"{horizon} dias — {item['priority']} · {item['finding']} · {item['action']} · owner {item['owner']}", styles["BodyText"]))
        else:
            story.append(Paragraph(f"{horizon} dias — nenhuma ação confirmada neste horizonte.", styles["BodyText"]))
    if executive["conditional_reviews"]:
        story.append(Paragraph("Validações antes de remediar", styles["Heading2"]))
        for item in executive["conditional_reviews"][:6]:
            story.append(Paragraph(f"{item['title']} — {item['reason']}", styles["BodyText"]))

    insights = data.get("discovery", {}).get("cross_domain_insights", [])
    story.append(Paragraph("Insights cruzados", styles["Heading2"]))
    for item in insights[:5]:
        story.append(Paragraph(f"{item.get('priority', 'P3')} — {item.get('title')} — responsável sugerido: {item.get('suggested_owner', 'Security & Governance')}", styles["BodyText"]))
        story.append(Spacer(1, 6))
    if context["limitations"]:
        story.append(Paragraph("Limitações da execução", styles["Heading2"]))
        for item in context["limitations"][:12]:
            story.append(Paragraph(f"{item['module']} — {item['status']} ({item['limitation_category']}): {item['likely_cause']} Próximo passo: {item['next_step']}", styles["BodyText"]))
    external = data.get("discovery", {}).get("external_assessments", {}).get("microsoft_zero_trust")
    if external:
        story.append(Spacer(1, 14))
        story.append(Paragraph("Microsoft Zero Trust Assessment — referência externa", styles["Heading2"]))
        story.append(Paragraph(f"Versão {external.get('source_version', 'N/D')} · {external.get('tests_total', 0)} verificações · score não combinado ao score do Assessment Engine.", styles["BodyText"]))
        status_line = " · ".join(f"{key}: {value}" for key, value in external.get("status_counts", {}).items())
        story.append(Paragraph(status_line, styles["BodyText"]))
        for pillar in external.get("pillars", []):
            pillar_status = " · ".join(f"{key}: {value}" for key, value in pillar.get("statuses", {}).items())
            story.append(Paragraph(f"{pillar.get('name')}: {pillar.get('tests')} verificações — {pillar_status}", styles["BodyText"]))

    discovery = data.get("discovery", {})
    graph = discovery.get("resource_map", {})
    hygiene = discovery.get("resource_hygiene", {})
    secure = discovery.get("defender_secure_score", {})
    story.append(Spacer(1, 14))
    story.append(Paragraph("Azure Intelligence", styles["Heading2"]))
    story.append(Paragraph(f"Resource Map: {graph.get('node_count', 0)} recursos e {graph.get('edge_count', 0)} relações demonstradas. Higiene: {hygiene.get('empty_resource_group_count', 0)} resource groups vazios, {hygiene.get('orphan_count', 0)} recursos órfãos/não associados e {hygiene.get('network_attention_count', 0)} sinais de rede em atenção.", styles["BodyText"]))
    for item in secure.get("top_improvements", [])[:5]:
        story.append(Paragraph(f"Secure Score — {item.get('control')}: ganho potencial {item.get('potential_score_increase')} ponto(s), {item.get('unhealthy_resources')} recursos não saudáveis.", styles["BodyText"]))
    document.build(story)


def write_one_page_brief(data: dict, path: Path, catalog: dict) -> None:
    """Gera briefing paisagem de uma página para abrir conversas com stakeholders."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    context = build_report_context(data)
    meta = data.get("metadata", {})
    profile = context["profile"]
    if profile in {"não informado", "unknown", ""} and meta.get("simulation", {}).get("is_simulation") is True:
        profile = "Demonstração sintética"
    quality = audit(data, catalog)["metrics"]
    logs = data.get("discovery", {}).get("collection_log", []) or []
    provisional = (quality.get("coverage", 0) < 100 or
                   any(item.get("status") in {"partial", "not_available", "error", "not_run"} for item in logs))
    score = quality.get("overall_score")
    score_label = f"{score:.1f}/100" if isinstance(score, (int, float)) else "N/D"
    findings = prioritize_findings(data.get("findings", []), data.get("metadata", {}).get("evidence_quality", {}), data.get("discovery", {}).get("lifecycle", {}).get("summary", {}).get("Custo mensal potencial", "Não quantificado"))
    executive = build_executive_intelligence(findings)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("BriefTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=colors.HexColor("#40205f"), spaceAfter=4)
    body_style = ParagraphStyle("BriefBody", parent=styles["BodyText"], fontSize=9, leading=12, spaceAfter=4)
    heading_style = ParagraphStyle("BriefHeading", parent=styles["Heading2"], fontSize=12, leading=15, textColor=colors.HexColor("#40205f"), spaceBefore=7, spaceAfter=5)
    small_style = ParagraphStyle("BriefSmall", parent=body_style, fontSize=8, leading=10)
    document = SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=30, rightMargin=30, topMargin=25, bottomMargin=25, title="Assessment — briefing de uma página", author="SoftwareOne Security & Governance Assessment")
    story = [Paragraph("Security &amp; Governance Assessment — briefing executivo", title_style)]
    story.append(Paragraph(f"{escape(str(meta.get('customer_name', 'Tenant')))} · Perfil {escape(str(profile))} · Execução {escape(str(context['started_at']))} · {escape(str(context['classification']))}", small_style))
    count_line = " · ".join(f"{escape(str(key))}: {value}" for key, value in context["status_counts"].items()) or "Sem registros de coletores"
    module_count = f"{context['status_counts'].get('success', 0)} concluídos · {context['status_counts'].get('partial', 0)} parciais · {context['status_counts'].get('not_available', 0) + context['status_counts'].get('error', 0)} indisponíveis/erro · {context['status_counts'].get('not_run', 0)} fora do perfil"
    coverage_value = quality.get("coverage")
    coverage_label = f"{coverage_value:.1f}% dos controles" if isinstance(coverage_value, (int, float)) else "N/D"
    score_cell = f"Score: {score_label}" + (" · PROVISÓRIO" if provisional else "")
    metrics = Table([[Paragraph(f"<b>{escape(score_cell)}</b>", body_style), Paragraph(f"<b>Cobertura:</b> {escape(coverage_label)}", body_style), Paragraph(f"<b>Achados:</b> {len(findings)}", body_style), Paragraph(f"<b>Duração:</b> {escape(str(context['duration_seconds'] or 'N/D'))} s", body_style)]], colWidths=[185, 165, 110, 130])
    metrics.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#f1f6fa")), ("BOX", (0,0), (-1,-1), .5, colors.HexColor("#cbd5e1")), ("INNERGRID", (0,0), (-1,-1), .3, colors.HexColor("#dbe3eb")), ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("LEFTPADDING", (0,0), (-1,-1), 8), ("RIGHTPADDING", (0,0), (-1,-1), 8), ("TOPPADDING", (0,0), (-1,-1), 7), ("BOTTOMPADDING", (0,0), (-1,-1), 7)]))
    story.extend([Spacer(1, 8), metrics, Paragraph("Saúde da coleta", heading_style), Paragraph(escape(module_count), body_style), Paragraph(f"<font size='7'>{count_line}</font>", small_style)])
    story.append(Paragraph("Riscos para discussão", heading_style))
    for item in findings[:3]:
        story.append(Paragraph(f"• <b>{escape(str(item.get('title', 'Achado')))}</b> — risco {escape(str(item.get('risk_score', 'N/D')))} · alcance {escape(str(item.get('affected', 'N/D')))} {escape(str(item.get('affected_unit', 'itens')))} · confiança {escape(str(item.get('evidence_confidence', 'não avaliada')))}", body_style))
    if not findings:
        story.append(Paragraph("Sem achados priorizados disponíveis nesta execução.", body_style))
    summary = executive["summary"]
    story.append(Paragraph("Decisão e próximos passos", heading_style))
    story.append(Paragraph(f"P1: {summary['p1']} · P2: {summary['p2']} · quick wins: {summary['quick_wins']} · validar evidência: {summary['conditional_review']}.", body_style))
    if executive["quick_wins"]:
        for item in executive["quick_wins"][:2]:
            story.append(Paragraph(f"• <b>{escape(str(item['title']))}</b> — {escape(str(item['priority']))} · esforço {escape(str(item['effort']))} · owner {escape(str(item['owner']))}", small_style))
    story.append(Paragraph("Limitações que afetam a leitura", heading_style))
    if context["limitations"]:
        for item in context["limitations"][:4]:
            story.append(Paragraph(f"• <b>{escape(str(item['module']))}</b> — {escape(str(item['status']))} / {escape(str(item['limitation_category']))}: {escape(str(item['likely_cause']))} Próximo passo: {escape(str(item['next_step']))}", small_style))
    else:
        story.append(Paragraph("Nenhuma limitação de coletor registrada. Confirme o escopo e os owners antes de conclusões.", body_style))
    story.extend([Spacer(1, 7), Paragraph("Uso em reunião: validar escopo, explicar evidências insuficientes, escolher owners e acordar próximos passos. Este briefing não aprova remediação.", small_style), Paragraph("CONFIDENCIAL · Somente leitura · Compartilhar apenas com participantes autorizados · Ausência de evidência não é conformidade.", small_style)])
    document.build(story)


if __name__ == "__main__":
    main()
