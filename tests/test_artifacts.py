import json
import sys
import unittest
import tempfile
import subprocess
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from ai_payload import SENSITIVE_VALUE_PATTERNS, build
from validate_artifacts import validate, validate_payload
from security_identity_intelligence import build as build_security_identity_intelligence


class ArtifactTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        """Gera fixtures temporárias para o teste ser independente de dist/.

        O pacote distribuído não precisa carregar artefatos antigos; o teste
        deve validar o pipeline atual em um checkout limpo.
        """
        cls._temporary = tempfile.TemporaryDirectory(prefix="assessment-artifacts-")
        cls.artifact_dir = Path(cls._temporary.name) / "dist"
        cls.artifact_dir.mkdir(parents=True, exist_ok=True)
        data = ROOT / "mock" / "assessment.json"
        subprocess.run([sys.executable, "src/generate_report.py", "--data", str(data), "--output", str(cls.artifact_dir / "assessment.html")], cwd=ROOT, check=True, capture_output=True, text=True)
        subprocess.run([sys.executable, "src/export_artifacts.py", "--data", str(data), "--output-dir", str(cls.artifact_dir)], cwd=ROOT, check=True, capture_output=True, text=True)
        cls.ai_payload = Path(cls._temporary.name) / "ai-payload.json"
        subprocess.run([sys.executable, "src/ai_payload.py", "--data", str(data), "--output", str(cls.ai_payload)], cwd=ROOT, check=True, capture_output=True, text=True)

    @classmethod
    def tearDownClass(cls):
        cls._temporary.cleanup()

    def test_html_is_self_contained(self):
        html = (self.artifact_dir / "assessment.html").read_text(encoding="utf-8")
        HTMLParser().feed(html)
        self.assertNotIn("<script src=", html)
        self.assertNotIn("http://", html)
        self.assertNotIn("https://", html)
        self.assertIn("FinOps e ciclo de vida", html)
        self.assertIn("Visão visual das políticas", html)
        self.assertIn("ca-card", html)
        self.assertIn("Exclusões identificadas", html)
        self.assertIn("Prioridade combina risco", html)
        self.assertIn("Custo potencial", html)
        self.assertIn('data-view-target="executive"', html)
        self.assertIn('data-view-target="technical"', html)
        self.assertIn('data-view-target="full"', html)
        self.assertIn("risk-matrix", html)
        self.assertIn("top-coverage-banner", html)
        self.assertIn("Limites desta leitura", html)
        self.assertIn("Causa provável", html)
        self.assertIn("Próximo passo", html)

    def test_spreadsheets_and_deck_open(self):
        workbook = load_workbook(self.artifact_dir / "assessment-action-plan.xlsx", read_only=True)
        self.assertIn("Plano de ação", workbook.sheetnames)
        self.assertGreaterEqual(workbook["Plano de ação"].max_row, 2)
        self.assertIn("Prioridades executivas", workbook.sheetnames)
        executive_rows = list(workbook["Prioridades executivas"].values)
        self.assertTrue(any(row[0] == "Sinais executivos adicionais" for row in executive_rows if row))
        self.assertTrue(any(row[0] == "FinOps · economia potencial (limite superior)" for row in executive_rows if row))
        workbook.close()
        deck = Presentation(self.artifact_dir / "assessment-executive-summary.pptx")
        self.assertGreaterEqual(len(deck.slides), 7)
        deck_text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text_frame"))
        self.assertIn("Decisão executiva — próximos passos", deck_text)
        self.assertIn("Roadmap executivo — 30 / 60 / 90 dias", deck_text)
        self.assertIn("Cobertura e limitações da execução", " ".join(shape.text for shape in deck.slides[3].shapes if shape.has_text_frame))
        brief = self.artifact_dir / "assessment-one-page-brief.pdf"
        self.assertTrue(brief.exists())
        self.assertEqual(len(PdfReader(str(brief)).pages), 1)

    def test_executive_kpis_are_consistent_across_html_pdf_pptx_and_xlsx(self):
        from report_metrics import build_executive_metrics
        metrics = build_executive_metrics(json.loads((ROOT / "mock/assessment.json").read_text(encoding="utf-8")))
        html = (self.artifact_dir / "assessment.html").read_text(encoding="utf-8")
        deck = Presentation(self.artifact_dir / "assessment-executive-summary.pptx")
        pptx_text = " ".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
        pdf_text = " ".join((page.extract_text() or "") for page in PdfReader(self.artifact_dir / "assessment-executive-summary.pdf").pages)
        workbook = load_workbook(self.artifact_dir / "assessment-action-plan.xlsx", read_only=True)
        kpi_rows = {row[0]: row[1:] for row in workbook["Indicadores executivos"].iter_rows(min_row=2, values_only=True)}
        workbook.close()

        def normalized(value):
            return " ".join(str(value).split()).casefold()

        for metric in metrics:
            self.assertIn(f'data-kpi-id="{metric["id"]}" data-source-status="{metric["source_status"]}"', html)
            self.assertIn(f'<b>{metric["display_value"]}</b>', html)
            needle = normalized(f'{metric["label"]}: {metric["display_value"]} · fonte {metric["source_status"]}')
            self.assertIn(needle, normalized(pdf_text), metric["id"])
            self.assertIn(needle, normalized(pptx_text), metric["id"])
            self.assertEqual(kpi_rows[metric["label"]][0], metric["display_value"])
            self.assertEqual(kpi_rows[metric["label"]][1], metric["source_status"])
            self.assertEqual(kpi_rows[metric["label"]][2], ", ".join(metric["source_modules"]))

    def test_action_plan_contains_cross_domain_insight_rows(self):
        source = json.loads((ROOT / "mock/assessment.json").read_text(encoding="utf-8"))
        source.setdefault("discovery", {})["cross_domain_insights"] = [{"id": "I-TEST", "title": "Insight de governança", "severity": "high", "risk": 80, "priority": "P2", "suggested_owner": "Cloud Governance", "effort_band": "Médio"}]
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "plan.xlsx"
            from export_artifacts import write_xlsx
            write_xlsx(source, path)
            workbook = load_workbook(path, read_only=True)
            rows = list(workbook["Plano de ação"].values)
            workbook.close()
        self.assertTrue(any(row[0] == "I-TEST" for row in rows[1:]))

    def test_executive_artifacts_use_engagement_metadata(self):
        source = json.loads((ROOT / "mock/assessment.json").read_text(encoding="utf-8"))
        source["metadata"].update({"customer_name": "Cliente Demo", "engagement_name": "Discovery Beta", "consultant_name": "Consultor SWO", "classification": "Confidencial"})
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)
            from export_artifacts import write_pptx, write_pdf
            write_pptx(source, output / "summary.pptx")
            write_pdf(source, output / "summary.pdf")
            deck = Presentation(output / "summary.pptx")
            deck_text = " ".join(shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text"))
            from pypdf import PdfReader
            pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(output / "summary.pdf").pages)
        self.assertIn("Discovery Beta", deck_text)
        self.assertIn("Cliente Demo", deck_text)
        self.assertIn("Cliente Demo", pdf_text)

    def test_unavailable_identity_evidence_stays_unknown_in_exports(self):
        source = json.loads((ROOT / "mock/assessment.json").read_text(encoding="utf-8"))
        discovery = source.setdefault("discovery", {})
        discovery["users"] = []
        discovery["identity_risks"] = []
        discovery["conditional_access"] = []
        log = [{"module": "Graph", "status": "error", "records": 0, "note": "synthetic auth failure"}]
        discovery["security_identity_intelligence"] = build_security_identity_intelligence(discovery, log)
        values = discovery["security_identity_intelligence"]["identity_posture"]
        self.assertIsNone(values["privileged_without_mfa"])
        self.assertEqual(values["evidence_status"], "error")
        from generate_report import render_executive_security_kpis
        discovery["collection_log"] = log
        html_kpis = render_executive_security_kpis(source)
        self.assertIn("Sem evidência", html_kpis)
        self.assertNotIn("<span>Sem MFA</span><b>0</b>", html_kpis)
        self.assertNotIn("<span>Privilegiados sem MFA</span><b>0</b>", html_kpis)
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)
            from export_artifacts import write_pptx, write_pdf, write_xlsx
            write_pptx(source, output / "summary.pptx")
            write_pdf(source, output / "summary.pdf")
            write_xlsx(source, output / "summary.xlsx")
            deck = Presentation(output / "summary.pptx")
            pptx_text = " ".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
            pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(output / "summary.pdf").pages)
            workbook = load_workbook(output / "summary.xlsx", read_only=True)
            xlsx_text = " ".join(str(cell.value) for row in workbook["Prioridades executivas"].iter_rows() for cell in row if cell.value is not None)
            workbook.close()
        for artifact_text in (pptx_text, pdf_text, xlsx_text):
            self.assertIn("N/D", artifact_text)
            self.assertNotIn("privilegiados sem MFA observado: 0", artifact_text)
            self.assertNotIn("privilegiados sem MFA observado: None", artifact_text)
            self.assertNotIn("economia realizável None", artifact_text)
        self.assertIn("fonte: error", pptx_text)
        self.assertIn("fonte error", pdf_text)

    def test_successful_empty_identity_collection_is_zero(self):
        result = build_security_identity_intelligence({}, [{"module": "Graph", "status": "success", "records": 0}])
        self.assertEqual(result["identity_posture"]["privileged_without_mfa"], 0)
        self.assertEqual(result["identity_posture"]["evidence_status"], "success")

    def test_ai_payload_contains_aggregates_only(self):
        source = json.loads((ROOT / "mock/assessment.json").read_text(encoding="utf-8"))
        serialized = json.dumps(build(source), ensure_ascii=False)
        for forbidden in ("contoso.example", "Ana Souza", "Bruno Lima", "vm-prd-web-01", "resource_id"):
            self.assertNotIn(forbidden, serialized)

    def test_ai_payload_contains_safe_azure_posture_aggregates(self):
        source = json.loads((ROOT / "mock/assessment.json").read_text(encoding="utf-8"))
        source.setdefault("discovery", {})["security_posture_summary"] = {
            "resources": 20,
            "resources_with_explicit_signals": 3,
            "signals_total": 4,
            "by_signal": [{"signal": "Storage: blob público habilitado", "resources": 2}],
        }
        result = build(source)
        self.assertEqual(result["azure_security_posture"]["resources_with_explicit_signals"], 3)
        self.assertEqual(result["azure_security_posture"]["signals"][0]["resources"], 2)
        self.assertNotIn("resource_id", json.dumps(result, ensure_ascii=False))

    def test_ai_payload_contains_safe_rbac_aggregates(self):
        source = json.loads((ROOT / "mock/assessment.json").read_text(encoding="utf-8"))
        source.setdefault("discovery", {})["rbac_summary"] = {
            "assignments": 12,
            "high_risk_assignments": 4,
            "critical_assignments": 2,
            "permanent_or_unknown_assignments": 10,
            "by_scope": [{"scope_kind": "Subscription", "assignments": 3}],
        }
        result = build(source)
        self.assertEqual(result["azure_rbac_posture"]["critical_assignments"], 2)
        self.assertEqual(result["azure_rbac_posture"]["by_scope"][0]["assignments"], 3)
        self.assertIn("não representa uso efetivo", result["azure_rbac_posture"]["interpretation"])
        self.assertNotIn("resource_id", json.dumps(result, ensure_ascii=False))

    def test_ai_payload_contains_safe_governance_aggregates(self):
        source = json.loads((ROOT / "mock/assessment.json").read_text(encoding="utf-8"))
        source.setdefault("discovery", {})["governance_summary"] = {"resources_assessed": 10, "without_owner": 3, "without_environment_tag": 4, "without_tags": 2, "policy_evaluated": 8, "policy_non_compliant": 1, "policy_compliance_rate": 87.5}
        result = build(source)
        self.assertEqual(result["azure_governance_posture"]["without_owner"], 3)
        self.assertEqual(result["azure_governance_posture"]["policy_compliance_rate"], 87.5)
        self.assertIn("não prova risco", result["azure_governance_posture"]["interpretation"])

    def test_artifact_validator_accepts_generated_package(self):
        result = validate(self.artifact_dir, self.ai_payload)
        self.assertEqual(result["status"], "valid", result["errors"])
        self.assertTrue(result["read_only"])

    def test_artifact_validator_blocks_invalid_contract(self):
        assessment = json.loads((ROOT / "mock/assessment.json").read_text(encoding="utf-8"))
        assessment["metadata"]["contract_status"] = "invalid"
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "assessment.json"
            path.write_text(json.dumps(assessment), encoding="utf-8")
            result = validate(self.artifact_dir, self.ai_payload, path)
        self.assertEqual(result["status"], "invalid")
        self.assertTrue(result["checks"]["contract"])

    def test_executive_artifacts_contain_no_sensitive_patterns(self):
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(self.artifact_dir / "assessment-executive-summary.pdf").pages)
        deck = Presentation(self.artifact_dir / "assessment-executive-summary.pptx")
        pptx_text = " ".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
        for pattern in SENSITIVE_VALUE_PATTERNS:
            self.assertIsNone(pattern.search(pdf_text), f"Padrão sensível no PDF: {pattern.pattern}")
            self.assertIsNone(pattern.search(pptx_text), f"Padrão sensível no PPTX: {pattern.pattern}")

    def test_ai_payload_validator_rejects_identity_fields(self):
        with tempfile.NamedTemporaryFile(mode="w", suffix=".json", encoding="utf-8") as handle:
            json.dump({"user_principal_name": "person@example.com"}, handle)
            handle.flush()
            path = Path(handle.name)
            self.assertTrue(validate_payload(path))


if __name__ == "__main__":
    unittest.main()
